from pathlib import Path
from datetime import datetime
from collections import Counter, defaultdict
import hashlib, json, re, sys, zipfile
import xml.etree.ElementTree as ET
from docx import Document
from pypdf import PdfReader

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'deliverables/G04-项目计划'
QA = ROOT / '.codex-build/revise-20261006'
NS = {'p':'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}

def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def compact(s):
    return re.sub(r'\s+', '', s)

deck = OUT / '城院帮项目计划G04.pptx'
slides = []
with zipfile.ZipFile(deck) as z:
    names = sorted((n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)),
                   key=lambda n:int(re.search(r'slide(\d+)', n).group(1)))
    assert len(names) == 48
    for num, name in enumerate(names, 1):
        xml = ET.fromstring(z.read(name))
        blocks = []
        for sp in xml.findall('.//p:sp', NS):
            paras = [''.join(t.text or '' for t in p.findall('.//a:t', NS))
                     for p in sp.findall('.//a:p', NS)]
            s = '\n'.join(paras).strip()
            if s:
                blocks.append(s)
        slides.append({'slide':num, 'blocks':blocks})

subs = json.loads((OUT/'图表与排期/子任务数据.json').read_text(encoding='utf-8-sig'))
idx = {q['code']:q for q in subs}
seen = defaultdict(list)
for sl in slides[21:43]:
    blocks = sl['blocks']
    k = blocks.index('编号') + 5
    while k + 4 < len(blocks) and re.fullmatch(r'\d+\.\d+\.\d+', blocks[k]):
        code, task, owner, pred, acceptance = blocks[k:k+5]
        q = idx[code]
        assert compact(owner).replace('（共同）', '') == compact(q['owner']), (sl['slide'], code, 'owner')
        assert set(re.findall(r'\d+\.\d+\.\d+', pred)) == set(q['preds']), (code, 'predecessors')
        parts = task.split('；')
        dates = re.findall(r'(\d{1,2})/(\d{1,2})', parts[1])
        assert dates, (code, 'dates')
        limits = [datetime.fromisoformat(q[v]).date() for v in ['start', 'finish']]
        actual_dates = []
        for m, d in dates:
            day = datetime(2027 if int(m) == 1 else 2026, int(m), int(d)).date()
            assert limits[0] <= day <= limits[1], (code, str(day), limits)
            actual_dates.append(day)
        full = re.search(r'全任务\s*([\d.]+)\s*日', task)
        duration = float(full.group(1)) if full else float(re.search(r'([\d.]+)\s*日', parts[2]).group(1))
        assert duration == q['duration'], (code, duration, q['duration'])
        money = re.search(r'(?:新增现金|新增|现金预算)\s*([\d.]+)\s*元', task)
        assert money and float(money.group(1)) == q['budget'], (code, 'budget')
        assert compact(q['evidence']) in compact(acceptance), (code, 'evidence')
        seen[code].append(actual_dates)
        k += 5
assert set(seen) == set(idx) and sum(map(len, seen.values())) == 98
for code, chunks in seen.items():
    ends = sorted(d for chunk in chunks for d in chunk)
    assert ends[0] == datetime.fromisoformat(idx[code]['start']).date()
    assert ends[-1] == datetime.fromisoformat(idx[code]['finish']).date()

all_text = '\n'.join('\n'.join(s['blocks']) for s in slides)
assert all(f'M{n:02d}' in all_text for n in range(1, 15))
assert '3 日' in slides[14]['blocks'] and '本周2日\n总计3日' in slides[14]['blocks']
assert all(t in '\n'.join(slides[18]['blocks']) for t in ['11/4', '11/25'])
assert '当前状态：尚未确认' in '\n'.join(slides[19]['blocks'])
assert sum('当前尚未确认' in s for s in slides[47]['blocks']) == 2
assert all(t in all_text for t in ['WorkBuddy', 'JMeter', '690.81', '100并发仅作容量探测'])

plan = Document(OUT/'城院帮项目计划G04.docx')
doc_text = '\n'.join(p.text for p in plan.paragraphs) + '\n' + '\n'.join(c.text for t in plan.tables for row in t.rows for c in row.cells)
assert len(plan.tables) > 30 and all(q['code'] in doc_text for q in subs)
assert '13.2工作量估算' in compact(doc_text) and '13.2成本估算' not in compact(doc_text)
assert '尚未确认' in doc_text
page_counts = {n:len(PdfReader(OUT/n).pages) for n in ['城院帮项目计划G04.pdf', '城院帮可行性分析报告G04.pdf']}
assert list(page_counts.values()) == [45, 26]
table_fonts = Counter(round(r.font.size.pt, 1) for t in plan.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.text and r.font.size)
assert set(table_fonts) == {10.5}, table_fonts
finalizer = json.loads((QA/'finalizer-latest.json').read_text(encoding='utf-8'))
assert digest(deck) in json.dumps(finalizer)
assert digest(ROOT/'deck6/exports/城院帮项目计划（G04）.pptx') == 'a5e48276ec8a39aaf77e7fd3cee419f73d626bd1d3c2629aa381c1c9eba7ba59'
result = {'slides':48, 'appendix_rows':98, 'unique_subtasks':95, 'weekly_management':14,
          'owners_dates_durations_predecessors_budgets_evidence':'matched',
          'pdf_pages':page_counts, 'table_font_pt':10.5, 'actual_meeting_and_evaluation':'unconfirmed',
          'source_ppt':'preserved', 'ppt_sha256':digest(deck)}
(QA/'release-text.json').write_text(json.dumps(slides, ensure_ascii=False, indent=2), encoding='utf-8')
(QA/'revision-check.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(result, ensure_ascii=False))
