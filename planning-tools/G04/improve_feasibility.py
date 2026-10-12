from pathlib import Path
import shutil, re, json, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont
from xml.sax.saxutils import escape
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT/'deliverables/G04-项目计划'
CH = OUT/'图表与排期'
HISTORY = ROOT/'archive/可行性分析历史/20261006'
HISTORY.mkdir(parents=True, exist_ok=True)
for name in ['城院帮可行性分析报告G04.md','城院帮可行性分析报告G04.docx','城院帮可行性分析报告G04.pdf','评审表16项对照.md','豆包PPT制作提示词.md','交付检查记录.md','请先阅读.md']:
    source=OUT/name
    if source.exists() and not (HISTORY/name).exists(): shutil.copy2(source,HISTORY/name)

FONT='C:/Windows/Fonts/msyh.ttc'

class Figure:
    def __init__(self,w,h,title):
        self.w,self.h=w,h
        self.im=Image.new('RGB',(w,h),'white');self.d=ImageDraw.Draw(self.im)
        self.svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}"><rect width="100%" height="100%" fill="white"/>']
        self.nodes=[];self.edges=[]
        self.text((w/2,40),title,42,True)
    def text(self,xy,text,size=36,center=True):
        font=ImageFont.truetype(FONT,size)
        lines=text.split('\n');height=size*1.35
        for i,l in enumerate(lines):
            box=self.d.textbbox((0,0),l,font=font);ww=box[2]-box[0]
            x=xy[0]-ww/2 if center else xy[0]
            y=xy[1]+i*height
            self.d.text((x,y),l,font=font,fill='#17202A')
            self.svg.append(f'<text x="{xy[0] if center else x}" y="{y+size}" text-anchor="{"middle" if center else "start"}" font-family="Microsoft YaHei" font-size="{size}" fill="#17202A">{escape(l)}</text>')
    def box(self,id,xywh,label,kind='entity'):
        x,y,w,h=xywh
        fill={'entity':'#F1F4F7','process':'#E3EFF4','store':'#FFFFFF','physical':'#EAF1F8'}[kind]
        if kind=='process':
            self.d.rounded_rectangle((x,y,x+w,y+h),radius=20,fill=fill,outline='#526373',width=3)
            self.svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="20" fill="{fill}" stroke="#526373" stroke-width="3"/>')
        else:
            self.d.rectangle((x,y,x+w,y+h),fill=fill,outline='#526373',width=3)
            self.svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{fill}" stroke="#526373" stroke-width="3"/>')
        if kind=='store':
            self.d.line((x+20,y,x+20,y+h),fill='#526373',width=3)
            self.svg.append(f'<line x1="{x+20}" y1="{y}" x2="{x+20}" y2="{y+h}" stroke="#526373" stroke-width="3"/>')
        size=38;lines=label.count('\n')+1
        self.text((x+w/2,y+(h-size*1.35*lines)/2),label,size)
        self.nodes.append((id,xywh,label,kind))
    def arrow(self,id,source,target,pts,label='',labelxy=None,size=34):
        self.d.line(pts,fill='#526373',width=3)
        x,y=pts[-1];xx,yy=pts[-2]
        import math
        a=math.atan2(y-yy,x-xx);s=14
        wings=[(x,y),(x-s*math.cos(a-.45),y-s*math.sin(a-.45)),(x-s*math.cos(a+.45),y-s*math.sin(a+.45))]
        self.d.polygon(wings,fill='#526373')
        self.svg.append('<polyline points="'+' '.join(f'{x},{y}' for x,y in pts)+'" fill="none" stroke="#526373" stroke-width="3"/>')
        self.svg.append('<polygon points="'+' '.join(f'{x},{y}' for x,y in wings)+'" fill="#526373"/>')
        if label: self.text(labelxy,label,size)
        self.edges.append((id,source,target,pts,label))
    def save(self,name):
        self.im.save(CH/(name+'.png'))
        (CH/(name+'.svg')).write_text('\n'.join(self.svg+['</svg>']),encoding='utf8')
        return self

f=Figure(1800,1020,'客户端与业务服务  独立运维与恢复')
f.box('c1',(40,170,420,170),'微信小程序\n读者与作者网页','physical')
f.box('c2',(40,450,420,140),'社区管理网页','physical')
f.box('b',(650,220,450,220),'Spring Boot\n认证 内容 互动 治理','physical')
f.box('db',(1360,210,380,130),'MySQL业务数据','store')
f.box('img',(1360,440,380,140),'正文图片\n存储目录','store')
f.box('opsweb',(40,760,420,140),'独立运维网页','physical')
f.box('exec',(650,750,450,160),'独立运维执行器\n固定目标与持久记录','physical')
f.box('bak',(1360,750,380,160),'版本包 备份\n日志与清单','store')
f.arrow('e1','c1','b',[(460,245),(650,245)],'请求', (555,190))
f.arrow('e2','b','c1',[(650,315),(460,315)],'结果',(555,325))
f.arrow('e3','c2','b',[(460,510),(565,510),(565,405),(650,405)],'管理数据',(555,525))
f.arrow('e4','b','db',[(1100,260),(1360,260)],'查询与保存',(1230,210))
f.arrow('e5','b','img',[(1100,385),(1230,385),(1230,490),(1360,490)],'上传与读取',(1240,525))
f.arrow('e6','opsweb','exec',[(460,815),(650,815)],'固定作业',(555,760))
f.arrow('e7','exec','opsweb',[(650,880),(460,880)],'状态结果',(555,900))
f.arrow('e8','exec','b',[(840,750),(840,440)],'部署与重启',(1010,615))
f.arrow('e9','exec','bak',[(1100,830),(1360,830)],'校验与保存',(1230,760))
f.arrow('e10','exec','db',[(1100,775),(1160,775),(1160,590),(1745,590),(1745,275),(1740,275)],'备份与恢复',(1470,625))
f.text((900,975),'运维入口独立于业务进程  备份同时包含数据库 图片与版本清单',30)
figures=[f.save('可行性系统流程图')]

f=Figure(1800,990,'城院帮与外部角色交换的数据')
f.box('u',(40,160,460,150),'E1 读者与认证作者')
f.box('a',(40,620,460,150),'E2 社区管理员')
f.box('p',(675,380,440,200),'0 城院帮平台','process')
f.box('mail',(1300,160,450,150),'E4 邮箱发送服务')
f.box('op',(1300,620,450,150),'E3 运维人员')
f.arrow('t1','u','p',[(500,205),(650,205),(650,420),(675,420)],'认证 检索 创作\n互动与反馈',(755,155),34)
f.arrow('t2','p','u',[(675,480),(570,480),(570,280),(500,280)],'文章 操作结果\n状态与允许查看的信息',(320,340),34)
f.arrow('t3','a','p',[(500,650),(600,650),(600,540),(675,540)],'审核与处置数据',(340,530),34)
f.arrow('t4','p','a',[(760,580),(760,835),(300,835),(300,770)],'举报 待审材料 处理结果',(625,860),34)
f.arrow('t5','p','mail',[(1115,420),(1210,420),(1210,205),(1300,205)],'认证邮件',(1440,355),34)
f.arrow('t6','mail','p',[(1300,280),(1155,280),(1155,480),(1115,480)],'发送结果',(1030,225),34)
f.arrow('t7','op','p',[(1300,650),(1200,650),(1200,540),(1115,540)],'运维作业请求',(1450,525),34)
f.arrow('t8','p','op',[(1045,580),(1045,835),(1530,835),(1530,770)],'状态 日志 备份清单',(1350,860),34)
f.text((900,945),'箭头只表示数据流  不表示操作顺序或工程任务依赖',30)
figures.append(f.save('可行性顶层数据流图'))

f=Figure(1800,1700,'一级业务处理与主要逻辑存储')
ys=[260,540,820,1170,1450]
inputs=['认证申请','检索与创作数据','互动数据','举报反馈','运维作业请求']
outputs=['认证状态','文章与操作结果','评论与操作结果','处理结果','状态 日志 备份']
procs=['1.0 身份认证','2.0 经验内容','3.0 评论与激励','4.0 社区治理','5.0 运维作业']
stores=['D1 用户与认证','D2 文章草稿与图片','D3 评论收藏与币事件','D4 举报反馈与处置','D5 运维记录与备份清单']
for i,y in enumerate(ys):
    ent='E1 读者与作者' if i<4 else 'E3 运维人员'
    f.box(f'u{i}',(30,y,350,135),ent)
    f.box(f'p{i}',(680,y,410,135),procs[i],'process')
    f.box(f'd{i}',(1300,y,470,135),stores[i],'store')
    f.arrow(f'i{i}',f'u{i}',f'p{i}',[(380,y+38),(680,y+38)],inputs[i],(530,y-5),33)
    f.arrow(f'o{i}',f'p{i}',f'u{i}',[(680,y+104),(380,y+104)],outputs[i],(530,y+115),33)
    f.arrow(f'w{i}',f'p{i}',f'd{i}',[(1090,y+38),(1300,y+38)],'保存数据',(1190,y-5),31)
    f.arrow(f'r{i}',f'd{i}',f'p{i}',[(1300,y+104),(1090,y+104)],'查询结果',(1190,y+115),31)
f.box('mail',(1310,150,450,65),'E4 邮箱发送服务')
f.arrow('mf','p0','mail',[(1010,260),(1010,175),(1310,175)],'认证邮件',(1130,125),30)
f.arrow('mr','mail','p0',[(1310,205),(1060,205),(1060,260)],'发送结果',(1180,216),30)
f.box('authadmin',(30,140,350,65),'E2 社区管理员')
f.arrow('authreview','authadmin','p0',[(380,160),(560,160),(560,275),(680,275)],'材料审核数据',(560,105),30)
f.arrow('authpending','p0','authadmin',[(680,290),(620,290),(620,195),(380,195)],'待审材料与结果',(520,207),30)
f.box('admin',(20,1060,410,65),'E2 社区管理员')
f.arrow('ar','admin','p3',[(430,1080),(610,1080),(610,1220),(680,1220)],'内容处置数据',(550,1015),31)
f.arrow('as','p3','admin',[(680,1185),(650,1185),(650,1115),(430,1115)],'举报与处理结果',(960,1105),31)
f.arrow('auth','d0','p1',[(1540,395),(1540,460),(880,460),(880,540)],'认证状态',(1450,415),31)
f.arrow('auth2','d0','p2',[(1540,395),(1790,395),(1790,735),(850,735),(850,820)],'认证状态',(1600,690),31)
f.arrow('article','d1','p2',[(1420,675),(1420,790),(1030,790),(1030,820)],'文章标识',(1410,747),30)
f.text((900,1640),'E1与E2各为同一外部实体的重复绘制  D1至D5为逻辑存储  备份独立保存',29)
figures.append(f.save('可行性一级数据流图'))

mx=ET.Element('mxfile',host='app.diagrams.net')
for index,fig in enumerate(figures):
    page=ET.SubElement(mx,'diagram',id=f'feas{index}',name=['系统处理流程','顶层数据流','一级数据流'][index])
    model=ET.SubElement(page,'mxGraphModel',pageWidth=str(fig.w),pageHeight=str(fig.h))
    root=ET.SubElement(model,'root');ET.SubElement(root,'mxCell',id='0');ET.SubElement(root,'mxCell',id='1',parent='0')
    for id,(x,y,w,h),lab,kind in fig.nodes:
        style='rounded=1;' if kind=='process' else 'rounded=0;'
        cell=ET.SubElement(root,'mxCell',id=id,value=lab,vertex='1',parent='1',style=style+'whiteSpace=wrap;html=0;fontSize=38;fontFamily=Microsoft YaHei;fillColor=#EAF1F8;strokeColor=#526373;')
        ET.SubElement(cell,'mxGeometry',x=str(x),y=str(y),width=str(w),height=str(h),attrib={'as':'geometry'})
    for id,source,target,pts,lab in fig.edges:
        cell=ET.SubElement(root,'mxCell',id=id,value=lab,edge='1',parent='1',source=source,target=target,style='edgeStyle=orthogonalEdgeStyle;endArrow=block;html=0;fontSize=31;fontFamily=Microsoft YaHei;')
        geo=ET.SubElement(cell,'mxGeometry',relative='1',attrib={'as':'geometry'});arr=ET.SubElement(geo,'Array',attrib={'as':'points'})
        for x,y in pts[1:-1]:ET.SubElement(arr,'mxPoint',x=str(x),y=str(y))
ET.ElementTree(mx).write(CH/'可行性模型图.drawio',encoding='utf-8',xml_declaration=True)

content=(Path(__file__).parent/'feasibility_content.md').read_text(encoding='utf8')
(OUT/'城院帮可行性分析报告G04.md').write_text(content,encoding='utf8')

def font(run,size=12,bold=False):
    run.font.name='Microsoft YaHei';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.left_margin=sec.right_margin=Cm(1.8);sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.7)
for name,size in [('Normal',12),('Title',22),('Subtitle',14),('Heading 1',16),('Heading 2',14),('Heading 3',12),('TOC 1',12)]:
    if name not in doc.styles:doc.styles.add_style(name,1)
    st=doc.styles[name];st.font.name='Microsoft YaHei';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
    st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.13
    if name.startswith('Heading'):st.paragraph_format.space_before=Pt(10);st.paragraph_format.keep_with_next=True
    if st._element.pPr is not None:
        for b in list(st._element.pPr.findall(qn('w:pBdr'))):st._element.pPr.remove(b)
doc.core_properties.author='G04';doc.core_properties.title='城院帮可行性分析报告'
doc.core_properties.subject='软件工程课程项目 可行性研究与计划阶段评审'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run('第 '),10.5);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld);font(footer.add_run(' 页'),10.5)
doc.add_paragraph('城院帮可行性分析报告','Title');doc.add_paragraph('软件工程课程项目 G04','Subtitle')
doc.add_paragraph('组长兼项目负责人 王晨 32401211\n成员 刘景浩 32401205')
doc.add_paragraph('2026年10月6日')
doc.add_paragraph('实习与求职经验分享社区\n免费阅读 非现金激励 自愿站外联系')
doc.add_paragraph('研究结论','Heading 2')
doc.add_paragraph('建议在限定范围、两人协作和本地课堂演示条件下有条件推进。报告说明方案取舍、技术验证、成本资源、进度组织及运营维护条件，并提供教师16项评审对照和验证登记清单。')
doc.add_paragraph('课程交付与长期运营分别判断。关键技术与作者供给的实际证据需要在对应阶段补充，不能以选型、图表或演示数据代替。')
doc.add_page_break();toc_title=doc.add_paragraph();font(toc_title.add_run('目录'),16,True)
tocp=doc.add_paragraph();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u');tocp._p.append(fld)
doc.add_page_break()

def table(headers,rows):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    if len(headers)==3:
        widths=[3.8,6.5,7.1] if headers[0]!='条目' else [1.2,6.8,9.4]
    elif len(headers)==4:widths=[3.2,4.5,4.7,5.0] if headers[0]=='方案' else [2.0,4.5,4.7,6.2]
    else:widths=[4,13.4]
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    border=OxmlElement('w:tblBorders')
    for side in ['top','bottom','left','right','insideH','insideV']:
        item=OxmlElement('w:'+side);item.set(qn('w:val'),'single');item.set(qn('w:sz'),'4');item.set(qn('w:color'),'D9D9D9');border.append(item)
    t._tbl.tblPr.append(border)
    for c,s in zip(t.rows[0].cells,headers):c.text=s
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        for c,s in zip(t.add_row().cells,values):c.text=s.replace('<br>','\n')
    for ri,row in enumerate(t.rows):
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for side in ['top','bottom','left','right']:
                item=OxmlElement('w:'+side);item.set(qn('w:val'),'single');item.set(qn('w:sz'),'4');item.set(qn('w:color'),'D9D9D9');b.append(item)
            pr.append(b)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DDE7EF' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            for p in c.paragraphs:
                p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.08
                p.paragraph_format.keep_with_next=(ri==0)
                if (headers[0]=='条目' or len(headers)==4) and ci==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:font(run,10.5,ri==0)
    gap=doc.add_paragraph();gap.paragraph_format.space_after=Pt(3);gap.paragraph_format.line_spacing=Pt(1);gap.paragraph_format.keep_with_next=True

lines=content.splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line or line.startswith('# '):i+=1;continue
    if line.startswith('## '):
        p=doc.add_paragraph(line[3:],'Heading 1')
        if line.startswith('## 附录') or line.startswith('## 12 '):p.paragraph_format.page_break_before=True
        i+=1;continue
    if line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2');i+=1;continue
    if line.startswith('|'):
        records=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            vals=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',x) for x in vals):records.append(vals)
            i+=1
        table(records[0],records[1:]);continue
    if line.startswith('!['):
        m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(OUT/m.group(2)),width=Cm(17.4))
        inline=p.runs[0]._element.xpath('.//wp:docPr')
        if inline:inline[0].set('descr',m.group(1))
        cap=doc.add_paragraph(m.group(1));cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in cap.runs:font(r,10.5)
        i+=1;continue
    doc.add_paragraph(line);i+=1
doc.save(OUT/'城院帮可行性分析报告G04.docx')

# Keep the teacher's full matrix as a companion, updating only the feasibility evidence.
p=OUT/'评审表16项对照.md';s=p.read_text(encoding='utf8')
s=s.replace('| 3 | 可行性与关键技术 | 可行性报告；2.1—2.3 | 方案与真实验证结果分开 |','| 3 | 可行性与关键技术 | 可行性报告第3至11节、5.2矩阵和附录C；2.1—2.3 | 技术路线、早期核查与完整验收分开；方案比较、成本效益、风险和推进条件齐全 |')
p.write_text(s,encoding='utf8')
p=OUT/'请先阅读.md';s=p.read_text(encoding='utf8')
s=s.replace('2. 城院帮可行性分析报告G04.pdf或Word：关键技术、资源、操作、进度与风险条件。','2. 城院帮可行性分析报告G04.pdf或Word：推荐与备选方案、系统流程和数据流、关键技术验证、成本效益、进度组织、操作维护及风险条件；附教师16项对照、内部审查清单和待执行验证登记。')
p.write_text(s,encoding='utf8')
p=OUT/'豆包PPT制作提示词.md';s=p.read_text(encoding='utf8')
start='## 可行性分析配套材料使用要求'
if start in s:s=s[:s.index(start)].rstrip()
s+='\n\n'+start+'\n\n可行性报告已按课程方法补充推荐方案与备选方案、成本效益、系统模型、关键技术验证和评审依据。风险可行性页呈现“有条件可行”、主要风险和10月15日内部结论确认；不要改写为关键技术全部验证通过。技术验证分早期条件核查与完整系统验收。当前现金预算沿用10月1日基准690.81元，4个月，其他新增费用0元。\n\n附录可按讲解需要增加1至2页：方案比较与继续推进条件、FV01至FV08验证和教师第3项证据。图表与排期中的可行性系统流程图.png、可行性顶层数据流图.png、可行性一级数据流图.png供选择，draw.io和SVG为可编辑源；一图一页，不缩成无法阅读的小图。完整16项对照见可行性报告附录A，内部清单是小组自查标准，不冒充教师新评审表。\n'
p.write_text(s,encoding='utf8')
print(json.dumps({'docx':str(OUT/'城院帮可行性分析报告G04.docx'),'tables':len(doc.tables),'model_figures':3,'teacher_items':16,'verification_items':8,'history':str(HISTORY)},ensure_ascii=False))
