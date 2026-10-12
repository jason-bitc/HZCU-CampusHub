from pathlib import Path
import json,re,csv,hashlib
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'deliverables/G04-项目计划'
docx=OUT/'城院帮可行性分析报告G04.docx'
pdf=docx.with_suffix('.pdf')
doc=Document(docx);reader=PdfReader(pdf)
assert len(reader.pages)==26
assert all(len(p.extract_text().strip())>100 for p in reader.pages)
assert doc.styles['Normal'].font.size.pt==12
assert all(abs(s.page_width.cm-21)<.01 and abs(s.page_height.cm-29.7)<.01 for s in doc.sections)
assert len(doc.inline_shapes)==3
assert len(doc.tables[15].rows)==17
assert [r.cells[0].text.strip() for r in doc.tables[15].rows[1:]]==[str(i) for i in range(1,17)]
assert len(doc.tables[-1].rows)==9
assert [r.cells[0].text.strip() for r in doc.tables[-1].rows[1:]]==[f'FV{i:02}' for i in range(1,9)]
for t in doc.tables:
    assert t.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
    for row in t.rows:
        assert row._tr.trPr.find(qn('w:cantSplit')) is not None
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    assert not r.text or r.font.size.pt==10.5
text='\n'.join(p.extract_text() for p in reader.pages)
for word in ['WorkBuddy','690.81','有条件','数据字典','关键技术','风险','附录A','附录B','附录C']:
    assert word in text.replace(' ','')
assert not re.search(r'目录\s*\.+\s*2',reader.pages[1].extract_text())
md=(OUT/'城院帮可行性分析报告G04.md').read_text(encoding='utf8')
assert all(f'FV{i:02}' in md for i in range(1,9))
assert all(x in md for x in ['原生Project','10月15日','10月18日22时','12月4日','2027年1月6日'])
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(OUT/'评审资料/计划阶段评审原表.pdf')==sha(Path('C:/Users/wangc/Downloads/软件工程-课程项目评审表-计划阶段评审.pdf'))
baseline={r['相对路径']:r['SHA256'] for r in csv.DictReader((OUT/'交付文件SHA256.csv').open(encoding='utf-8-sig'))}
stable=['城院帮项目计划G04.docx','城院帮项目计划G04.pdf','详细周计划G04.md','图表与排期/任务数据.json','图表与排期/子任务数据.json','图表与排期/周计划数据.json']
for name in stable:assert sha(OUT/name)==baseline[name],name
native=json.loads((ROOT/'.codex-build/feasibility-20261006/native_audit_result.json').read_text(encoding='utf8'))
assert native['count']==214 and native['semantic_diff_count']==0 and native['readonly_no_write']
assert sha(OUT/'图表与排期/城院帮项目计划G04.mpp')==native['mpp_sha256']
assert all(r['Overallocated']=='False' for r in native['resources'])
record=OUT/'交付检查记录.md'
s=record.read_text(encoding='utf8').replace('可行性报告8页','可行性报告26页')
heading='## 2026年10月6日可行性研究检查'
if heading in s:s=s[:s.index(heading)].rstrip()
s+='\n\n'+heading+'\n\n可行性报告按所提供第1章、第13章课件及2页计划阶段评审原表优化；教材方法按出版社公开目录核对，不声称已经取得整书正文。报告为26页、12节正文和3项附录，包含推荐与备选方案、系统流程图、顶层及一级数据流图、简要数据字典、8项关键技术验证、成本效益、资源进度、操作维护、风险与推进条件。\n\n教师16项逐项核对到报告或配套证据；原表原样保存在评审资料。内部审查清单与教师表明确区分，产品验证登记仍为待验证，会议、成员评价和真实供给不代填结果。\n\nWord已更新目录并导出同版PDF；26页全部经图片检查，修正目录自引用、表格列宽与图表标签分页。A4、正文12pt、表格10.5pt、目录12pt，表头重复且行不跨页拆分。文档包检验不等于产品技术或系统测试已通过。\n\n项目计划Word/PDF、详细周计划和任务数据通过原文件摘要比对，保持一致；现有MPP经过只读核查，214项的名称、日期、工期、负责人、前置、总时差和成本与原任务核查表一致，两位资源均无过载。本次未保存或修改MPP。预算690.81元、4个月、原排期和分工未改变。豆包PPT提示词与评审对照已同步，图表提供PNG、SVG及draw.io源。原可行性材料归档在archive/可行性分析历史/20261006，交付恢复资料、文件摘要及压缩包随本次更新。\n'
record.write_text(s,encoding='utf8')
p=OUT/'请先阅读.md';s=p.read_text(encoding='utf8')
line='\n可行性报告为26页，附录A提供教师16项证据对照，附录B为小组内部审查，附录C为待执行验证登记；计划阶段评审原表在评审资料目录中保留原件。\n'
if line.strip() not in s:s+=line
p.write_text(s,encoding='utf8')
p=ROOT/'planning-tools/G04/README.md';s=p.read_text(encoding='utf8')
heading='## 可行性报告维护'
if heading in s:s=s[:s.index(heading)].rstrip()
s+='\n\n'+heading+'\n\n可行性正文源为feasibility_content.md。完整项目计划生成后，先执行improve_feasibility.py以保留研究内容和3张模型图，再执行export_feasibility.ps1更新Word目录和导出PDF。使用本机Word导出并将PDF逐页渲染检查；seal_feasibility.py核对教师16项、技术8项、版式及原计划排期未被改变，最后执行verify_and_package.py更新恢复资料与交付包。当前报告26页，其他计划文件不因本次内容优化重排。\n'
p.write_text(s,encoding='utf8')
result={'pages':26,'figures':3,'teacher_criteria':16,'technical_verifications':8,'tables':len(doc.tables),'typography':'A4/12pt body/10.5pt tables/12pt TOC','original_rubric':'hash matched','existing_plan_schedule':'unchanged','product_verifications':'pending, not claimed passed'}
(ROOT/'.codex-build/feasibility-20261006/validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(result,ensure_ascii=False))
