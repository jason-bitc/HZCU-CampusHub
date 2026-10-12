from pathlib import Path
from datetime import datetime
from decimal import Decimal
from collections import defaultdict
import hashlib, json, re, sys, zipfile
import xml.etree.ElementTree as ET
from docx import Document
from pypdf import PdfReader

sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'deliverables/G04-项目计划'
QA=ROOT/'.codex-build/project-report-20261008'
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
slides=[]
deck=OUT/'城院帮项目汇报G04.pptx'
with zipfile.ZipFile(deck) as z:
    assert len([n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)])==58
    for n in range(1,59):
        xml=ET.fromstring(z.read(f'ppt/slides/slide{n}.xml'))
        blocks=[]
        for sp in xml.findall('.//p:sp',NS):
            s='\n'.join(''.join(t.text or '' for t in pa.findall('.//a:t',NS)) for pa in sp.findall('.//a:p',NS)).strip()
            if s:blocks.append(s)
        tables=[[[ ''.join(t.text or '' for t in c.findall('.//a:t',NS)) for c in row.findall('a:tc',NS)] for row in tb.findall('a:tr',NS)] for tb in xml.findall('.//a:tbl',NS)]
        slides.append({'slide':n,'blocks':blocks,'tables':tables})

def compact(s):return re.sub(r'\s+','',s)
subs=json.loads((OUT/'图表与排期/子任务数据.json').read_text(encoding='utf-8-sig'))
idx={q['code']:q for q in subs};seen=defaultdict(list)
for sl in slides[31:53]:
    b=sl['blocks'];k=b.index('编号')+5
    while k+4<len(b) and re.fullmatch(r'\d+\.\d+\.\d+',b[k]):
        code,task,owner,pred,accept=b[k:k+5];q=idx[code]
        assert compact(owner).replace('（共同）','')==compact(q['owner']),(code,'owner')
        assert set(re.findall(r'\d+\.\d+\.\d+',pred))==set(q['preds']),(code,'pred')
        parts=task.split('；');dates=re.findall(r'(\d{1,2})/(\d{1,2})',parts[1])
        assert dates
        ds=[datetime(2027 if int(m)==1 else 2026,int(m),int(d)).date() for m,d in dates]
        lo,hi=[datetime.fromisoformat(q[v]).date() for v in ['start','finish']]
        assert all(lo<=d<=hi for d in ds),(code,'date')
        full=re.search(r'全任务\s*([\d.]+)\s*日',task)
        duration=float(full.group(1)) if full else float(re.search(r'([\d.]+)\s*日',parts[2]).group(1))
        assert duration==q['duration'],(code,'duration')
        money=re.search(r'(?:新增现金|新增|现金预算)\s*([\d.]+)\s*元',task)
        assert money and float(money.group(1))==q['budget'],(code,'budget')
        assert compact(q['evidence']) in compact(accept),(code,'evidence')
        seen[code].extend(ds);k+=5
assert len(seen)==95 and set(seen)==set(idx)
for code,ds in seen.items():
    assert min(ds)==datetime.fromisoformat(idx[code]['start']).date()
    assert max(ds)==datetime.fromisoformat(idx[code]['finish']).date()
alltext='\n'.join('\n'.join(s['blocks'])+'\n'+str(s['tables']) for s in slides)
assert all(f'M{n:02d}' in alltext for n in range(1,15))
for n in [4,11,12,13,14]:assert len(slides[n-1]['tables'])==1
budget=slides[12]['tables'][0]
assert sum(Decimal(budget[r][2]) for r in [1,2,3])==Decimal(budget[4][2])==Decimal('690.81')
assert all(f'{n:02d} / 58' in slides[n-1]['blocks'] for n in range(1,59))
assert '城院帮项目汇报' in slides[0]['blocks']
assert '可行性报告§3–11、附录C；本PPT第11–15页' in slides[56]['blocks']
assert sum('当前尚未确认' in s for s in slides[57]['blocks'])==2
assert len(PdfReader(OUT/'城院帮项目计划G04.pdf').pages)==45
assert len(PdfReader(OUT/'城院帮可行性分析报告G04.pdf').pages)==26
d=Document(OUT/'城院帮项目计划G04.docx')
doc_text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
assert '实际项目汇报PPT与配套图表' in doc_text
assert all(q['code'] in doc_text for q in subs)

header='''# 城院帮完整项目汇报PPT制作提示词

请制作软件工程课程G04小组的完整项目汇报PPT，文件名为城院帮项目汇报G04.pptx。汇报覆盖项目介绍、可行性分析、实施计划、质量维护和项目管理。封面与目录必须体现完整范围。

小组为G04。王晨32401211，组长兼项目负责人；刘景浩32401205，成员。王晨承担项目管理，两人共同参与开发、文档、设计评审和验收。AI辅助不替代学生的责任，不预设贡献比例。

## 产品及事实边界

城院帮是本校实习与求职经验分享社区，重点面向已知道目标岗位、正在准备投递和面试的大三大四学生。平台结构支持全校，首批演示集中在计算机相关专业。按岗位找经验，公司可选，同专业筛选保留。

免费阅读，认证后评论、点赞、非现金投币及收藏，作者自愿公开站外联系方式。认证用户可分享亲历，重点邀请已有实习或工作经历的人。网页Markdown创作，小程序基础图文与草稿，资料通过网盘链接说明来源。校园邮箱自动认证，必要时材料人工审核。社区先发布后管理，独立运维网页直接执行固定部署、重启和备份恢复。

不增加付费文章、现金支付、站内咨询订单、即时通讯、跑腿、悬赏任务或课程学习专区。前期校友圈截图与讨论支持问题描述，不能编造访谈数量、用户统计、作者承诺或增长率。认证不等于内容已经真实可靠。

可行性结论为限定范围及本地课堂演示条件下有条件可行。产品技术验证仍待按真实版本执行，长期作者供给仍需验证。实际会议与当前成员评价的最新人工确认是2026年10月6日尚未确认，保留待填表，不虚构日期、成果或比例。

正式环境为个人电脑Windows与WSL2，本地网页及微信开发者工具演示。腾讯云2核2GB仅为可选验证资源。预算按各1份、4个月：Plus538.81元，豆包152.00元，共690.81元。沿用2026年10月1日预算基准，80美元乘6.7351，不用舍入后的单月人民币金额重复相乘。其他新增费用0元，共用预算在2.2.1登记一次。不得展示密码或认证材料。

## 汇报结构和排版

总计58页。正文第1至31页，附录第32至58页。正文按下面逐页内容组织，详细周计划作为电子附录保留，课堂不要求逐页讲附录。

沿用当前白底、青蓝紫配色和Noto Sans SC字体。正文使用可投影阅读的字号，先减少重复措辞再调整字号。方案、技术验证和预算等表格使用可编辑对象。网络图完整等比例显示，不裁切标题、事件和图例。两位成员颜色同样醒目。

所有定位明确写项目计划报告、可行性报告或PPT页码。引用编号与讲者备注一致。课件条目只列名称、章节及课程来源，教材保留规范书目信息，未知版次明确待核对。不用旧课程学习或付费平台原型截图，不将示意场景冒充已运行界面。

## 正文逐页内容

'''
prompt=[header]
for sl in slides[:31]:
    blocks=[x for x in sl['blocks'] if x!='城院帮项目汇报 G04' and not re.fullmatch(r'\d{2} / 58',x)]
    title='城院帮项目汇报' if sl['slide']==1 else blocks[0]
    prompt.append(f"### 第{sl['slide']}页 {title}\n\n")
    for x in blocks[1:]:prompt.append(x+'\n\n')
    for tb in sl['tables']:
        prompt.append('| '+' | '.join(tb[0])+' |\n')
        prompt.append('| '+' | '.join(['---']*len(tb[0]))+' |\n')
        for row in tb[1:]:prompt.append('| '+' | '.join(v.replace('\n','<br>') for v in row)+' |\n')
        prompt.append('\n')
prompt.append('## 附录安排\n\n第32至53页完整周计划，第54至56页补充工程网络，第57至58页计划阶段16项评审对照。95项子任务及14项周管理完整保留，日期、任务负责人、前置、工期、预算和完成证据与报告及MPP一致。5.9工期3日，5.10跨周共3日；不把本周部分工期写成总工期。\n\n')
prompt.append((OUT/'详细周计划G04.md').read_text(encoding='utf-8')+'\n\n')
prompt.append((OUT/'评审表16项对照.md').read_text(encoding='utf-8')+'\n\n')
prompt.append('## 提交前核对\n\n标题、目录、背景、选题、目标用户、各用户核心功能、可行性、架构、工具、分工、排期、测试维护、参考资料与评价入口完整。报告日期和金额沿用任务数据，未知事实仍待确认。生成真实PPT文件后逐页检查，不以提示词代替汇报附件。\n')
(OUT/'豆包PPT制作提示词.md').write_text(''.join(prompt),encoding='utf-8')

readme=OUT/'请先阅读.md';s=readme.read_text(encoding='utf-8')
s=s.replace('城院帮项目计划G04.pptx：48页正式汇报材料，含项目介绍、排期、网络图、详细周计划附录及评审对照。','城院帮项目汇报G04.pptx：58页完整项目汇报。正文31页覆盖项目介绍、可行性、实施计划与维护，附录27页保留完整周计划、网络图及评审对照。')
s=s.replace('PPT已经与报告、排期及评审对照同步检查。','完整项目汇报PPT已经与报告、排期及评审对照同步检查。')
readme.write_text(s,encoding='utf-8')
entry=ROOT/'docs/PROJECT_PLAN.md';s=entry.read_text(encoding='utf-8').replace('城院帮项目计划G04.pptx','城院帮项目汇报G04.pptx').replace('项目计划PPT','完整项目汇报PPT')
s=s.replace('实际PPT包含48页，计划报告45页。','10月6日的计划汇报为48页，计划报告45页。')
s=re.sub(r'\n2026年10月8日将PPT组织为完整项目汇报[^\n]*\n?', '\n', s)
s+='\n2026年10月8日将PPT组织为完整项目汇报，正文31页、附录27页，新增项目背景、现状观察、目标、典型场景、作者激励及分项可行性说明。计划排期、95项子任务和预算保持一致。\n'
entry.write_text(s,encoding='utf-8')
check=OUT/'交付检查记录.md';s=check.read_text(encoding='utf-8')
s=s.replace('本目录已包含48页实际PPT，与报告及周计划附录同步检查。','本目录已包含58页完整项目汇报PPT，与报告及周计划附录同步检查。')
s=s.replace('实际PPT保留48页及95项子任务附录。','10月6日修订时，实际PPT保留48页及95项子任务附录。')
if '## 2026年10月8日完整项目汇报调整' in s:
    s=s[:s.index('## 2026年10月8日完整项目汇报调整')].rstrip()
s+='''

## 2026年10月8日完整项目汇报调整

PPT覆盖项目介绍、可行性分析、实施计划、质量维护和项目管理。正文31页，附录27页。项目介绍增加背景与现状、总体目标、读者场景和作者激励；可行性按方案比较、技术验证、经济资源、操作与持续运营及条件结论展开。保留原48个工作包、95项子任务、14项周管理、预算、排期和验收边界。

报告、PPT、提示词和评审定位同步。计划Word与PDF只更新实际项目汇报的引用，仍为45页；仅第43页发生显示变化，已复查，其余44页逐像素一致。可行性报告仍26页。

PPT已在本机PowerPoint打开并完整渲染，新增五个表格为可编辑对象，预算合计核查一致。正式交付前逐页检查正文和附录，评审第15、16项保留尚未确认。文档版本恢复资料、文件摘要和压缩包在最终材料确认后重新生成。

此前专门计划汇报及调整前材料归档在archive/项目汇报调整前/20261008。当前汇报入口为城院帮项目汇报G04.pptx，正式目录不并列两套汇报。
'''
check.write_text(s,encoding='utf-8')
(QA/'slides.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2),encoding='utf-8')
(QA/'content-check.json').write_text(json.dumps({'slides':58,'body':31,'appendix':27,'unique_subtasks':95,'tasks':'matched','budget':'690.81','meeting_evaluation':'unconfirmed','sha256':hashlib.sha256(deck.read_bytes()).hexdigest()},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'slides':58,'body':31,'appendix':27,'unique_subtasks':95,'task_fields':'matched','budget':'690.81','prompt':'synced','meeting_evaluation':'unconfirmed'},ensure_ascii=False))
