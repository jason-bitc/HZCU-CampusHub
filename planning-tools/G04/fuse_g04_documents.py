from pathlib import Path
from datetime import datetime,date,timedelta
import json,re
from docx import Document
from docx.shared import Cm,Pt,Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(r'C:\Users\wangc\Desktop\chengyuanbang-mp');OUT=ROOT/'deliverables/G04-项目计划';CH=OUT/'图表与排期'
first=(ROOT/'archive/项目计划历史/20261001/deck4/round4/materials/01_项目计划.md').read_text(encoding='utf8')
second=(ROOT/'archive/项目计划历史/20261001/deliverables/G04-project-plan-20261001/城院帮项目计划G04.md').read_text(encoding='utf8')
tasks=json.loads((CH/'任务数据.json').read_text(encoding='utf8'));subs=json.loads((CH/'子任务数据.json').read_text(encoding='utf8'));weeks=json.loads((CH/'周计划数据.json').read_text(encoding='utf8'))
idx={q['code']:q for q in subs}
def section_first(n):
 m=re.search(r'^## '+str(n)+r' [^\n]+\n',first,re.M);end=re.search(r'^## \d+ |^# 附录',first[m.end():],re.M)
 return first[m.end():m.end()+end.start() if end else len(first)].strip()
def section_second(n):
 m=re.search(r'^'+str(n)+r' [^\n]+\n',second,re.M);end=re.search(r'^\d+ [^\n]+\n',second[m.end():],re.M)
 return second[m.end():m.end()+end.start() if end else len(second)].strip()
def table(headers,rows):return '| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(str(v).replace('\n','；').replace('|','／') for v in row)+' |' for row in rows)+'\n'
def dt(s):return datetime.fromisoformat(s).strftime('%m/%d')
def datewindow(q):return dt(q['start'])+('—'+dt(q['finish']) if dt(q['start'])!=dt(q['finish']) else '')
def owner(q):return q['owner']+('（共同）' if '、' in q['owner'] else '')
def cash(q):return f'{q["duration"]:g}日；现金预算{q["budget"]:g}元'
refs='''[R1] 软件工程课程课件 第13章 软件项目管理，97页。依据13.1软件规模、13.2成本估算、13.3进度计划、13.4人员组织、13.5质量保证及13.6配置管理。

[R2] 软件工程课程课件 第1章 软件工程学概述，78页。依据软件生命周期、瀑布模型、文档及阶段评审。

[R3] 软件工程课程计划与团队作业通知，2026学年。课程节点、阶段成果及周日22时邮件提交要求。

[R4] 软件工程课程项目计划阶段评审检查表，2页，16项评审要求。

[R5] 云朵朵与校友圈界面观察截图、G04项目需求讨论记录，2026年9月。用于需求来源与初步竞品观察。

[R6] 张海藩、牟永敏 软件工程导论 第6版，清华大学出版社，2013，ISBN 9787302330981。教材结构参考；完整正文不在本次提供资料范围。

[R7] Kathy Schwalbe IT项目管理，机械工业出版社。版次待课程教材核对，项目管理结构参考，不引用未经核对的正文页码。

[R9] OpenAI官方文档 Pricing，https://learn.chatgpt.com/docs/pricing ，2026年10月1日核对；ChatGPT Plus每月20美元。

[R10] 中国银行 外汇牌价，https://www.boc.cn/sourcedb/whpj/ ，2026年10月1日17:59:27；美元中行折算价673.51元/100美元。

[R8] 工具官方资料 Microsoft Project 支持文档、Apache JMeter用户手册、JUnit用户指南、Git文档。用于工具类别、计划排期和验证方法。

课件参考条目按资料名称、章节与课程来源列示。项目管理方法以所提供课件和评审表为直接依据，任务估算与技术方案由本组制定。'''
intro=section_first(1)
intro=re.sub(r'项目周期：.*?计划覆盖需求、设计、开发、验证、文档、运营维护与收尾。','计划基准日期为2026年10月1日，执行排期从10月8日至2027年1月6日。王晨担任组长兼项目负责人，刘景浩担任组员。已过的问题定义和选题不再列为待办，本计划安排项目计划、可行性、需求、设计、实现、测试、部署维护及收尾。',intro)
intro=intro.replace('[R5]','[R5]')
scope=section_first(2)
life='''采用阶段成果与评审控制的开发过程，WBS按软件工程阶段组织。需求和总体、详细设计分别交付报告；11月5日完成内部设计评审后，11月6日开始集中实现。课堂在11月11日、18日和25日分别检查设计材料，新增意见通过影响分析进入模块调整，不将编码开始等同于课堂评审已通过。[R1][R2]

'''+table(['阶段','本项目工作与产物','阶段完成条件'],[
('项目计划与评审','WBS、周计划、Project、预算、风险与责任','每项任务可分配、可交接、可验收；评审意见有记录'),
('可行性研究','技术、经济、操作、进度、资料与风险分析','关键技术有验证方法，失败有替代方案'),
('需求分析','SRS、用例、原型、权限、数据与异常','需求有编号和明确验收条件；双方确认内部基线'),
('总体与详细设计','架构、ER、接口、流程、状态及运行恢复设计','关键问题关闭；文档足以指导实现和测试'),
('编码与单元测试','双端、业务和管理运维模块；单元与联调记录','AI输出人工核验；真实接口联通，形成候选版本'),
('综合测试','功能、接口、安全兼容、性能压力、稳定性与恢复证据','核心用例通过；阻塞和严重缺陷关闭；限制明确'),
('部署与维护','本机发布包、操作规程、维护变更与交接','可冷启动、恢复、回退；双方可操作管理流程'),
('总结与收尾','偏差分析、总结、真实评价、版本和最终交付','资料同版本，课程意见完成处理，交接有记录')])
tools='''客户端包括微信小程序、读者作者网页、社区管理网页和独立运维入口。uni-app组织小程序与读者作者H5；管理网页按社区与运维职责区分。Spring Boot业务单体提供共享认证、内容、互动和治理接口；MySQL保存业务数据。独立运维执行器在WSL2中按固定目标执行任务，业务服务停止时仍可查运维状态。

![系统架构与本机部署](图表与排期/架构与部署.png)

'''+table(['工程活动','采用工具','主要产物'],[
('计划与正式文档','Microsoft Word、Excel','计划、SRS、可行性、测试维护报告与追踪表'),('项目排期','Microsoft Project','原生MPP、阶段与详细甘特图、依赖、资源和关键任务'),('分析与设计','draw.io、Axure RP','架构、工程网络、ER、流程和交互原型；许可不足时用可用的同类原型工具'),('AI辅助开发','Codex、豆包、WorkBuddy','Codex搭框架和接口骨架；豆包辅助代码细节；小程序在WorkBuddy中开发；双方审查后入Git'),('前端','HBuilderX、微信开发者工具','uni-app构建与双端调试、开发者工具演示'),('后端与数据库','IntelliJ IDEA、Java、Spring Boot、MySQL','接口代码、SQL、初始化和数据库版本'),('单元接口及集成','JUnit、Spring测试支持、Postman','单元结果、接口集合、数据断言和集成证据'),('性能压力与稳定性','Apache JMeter','脚本、JTL、P95、错误率和资源趋势'),('缺陷跟踪','Bugzilla','复现、等级、指派、修复、复测和关闭记录'),('配置与协作','Git','代码文档图表、变更、提交历史及基线标签')])+'''
工具以覆盖工程活动为原则，不要求同时采用课件所举同类工具的全部产品。JMeter度量HTTP接口，不代替网页渲染和小程序体验检查。AI是辅助工具，两位学生对各自任务与最终证据负责。[R1][R8]'''
wbs='''WBS按8个工程阶段组织，48个工作包细化为95项子任务。工作包用于范围与交付管理，子任务用于逐周执行和接口交接。另列14项持续周管理活动、35项交叉复核及12个课程里程碑；这些活动进入Project资源检查，不充当软件功能点。

软件规模按课件的输入、输出、查询、主逻辑文件和外部接口做初步估算。外部输入18项×4、输出4项×5、查询8项×4、主逻辑文件9项×10、外部接口1项×7，共221个未调整点。该估计仅用于范围复核，不是严格功能点审计，也不换算人时；需求细化后按实际数据更新。[R1，13.1至13.2]

![阶段与成员任务总览](图表与排期/阶段甘特总览.png)

计划采用周一至周五排期日历，国庆10月1日至7日及元旦1月1日不排正常任务。工期是依赖与进度刻度，不限定成员固定工作时段。周末保留邮件提交与必要缓冲。全程排期跨度为64个工作日，不代表两人的总投入。

附录A给出48个工作包的时间、责任、依赖、产物与完成标准；附录B逐周列出具体子任务、交接点和验收证据。共用订阅预算690.81元在任务2.2.1登记一次；其余任务不重复分摊订阅费用。正文图展示主要阶段；完整子任务、复核和管理任务在原生MPP与分页甘特图中查阅。

工程网络与关键控制链

工程网络采用箭线表示法，箭头为活动，圆为事件；圆左为事件号，右上EET、右下LET；虚活动持续期0，仅表示依赖。图按分析设计、实现测试、维护收尾三个范围分页，跨页同号事件一致。[R1，13.3]

事件按工作日历计算：最早时刻取计划启动约束与所有前置完成时刻的最大值；最迟时刻从项目终点及内部截止倒推。活动总时差为终点LET减起点EET再减持续期。48个工作包的总时差已逐项与Microsoft Project核对。完整计算和日期约束见图表附件，不将纯逻辑持续期误作实际日历跨度。

'''+table(['控制链或节点','安排及管理含义'],[
('需求形成','1.1→1.2→2.2→2.3→3.1→3.2或3.3→3.4；保护10月23日内部需求基线'),('设计转实现','3.5→4.5；11月4日课后处理反馈，11月5日确认内部设计'),('实现与验收','5.4→5.5→5.9→5.10→5.11→4.7→5.12后分功能安全和性能压力两路，经6.6汇入6.7'),('后端与恢复支路','后端移动支路有1日总时差；恢复验证6.5有0.5日；不能据此省略验证'),('测试维护与收尾','报告复核、指南、冷启动、交接、总结及归档设置明确内部截止；1月6日为固定课程节点')])+'''
5.4至5.11的网页、社区和运维串行安排包含刘景浩的资源顺序，不能全部理解为技术强依赖。组长管理和交叉复核单独计入负荷。任务延期先分析交接、截止和人员冲突，再调整子任务，不以取消测试补回时间。'''
dates=[('10/14','10/11','项目计划评审','1.2评审材料；1.3课后修订'),('10/21','10/18','可行性报告','2.3内部结论与交叉确认'),('10/28','10/25','SRS与需求PPT','3.4需求基线与完整用例'),('11/04','11/01','SRS修订与需求复审','先提交已有修订；3.5处理当堂新意见'),('11/11','11/08','总体设计报告','4.5内部设计基线'),('11/18','11/15','详细设计报告','设计正文、接口和状态一致'),('11/25','11/22','设计修订与翻转PPT1','先交现有修订；4.7处理当堂意见'),('12/02','11/29','代码清单、用例与翻转PPT2','5.12联调与候选版本'),('12/09','12/06','系统测试报告','6.7首份报告；6.8后续复核'),('12/16','12/13','测试报告修订','先交已有修订；7.5处理当堂意见'),('12/30','12/27','总结与总评PPT','8.3交付核查与演练'),('01/06','01/03','归档基线与课程收尾','12/31准备提交基线；01/04核对尾项；01/06确认')]
milestones='''课堂节点依据课程计划，邮件在此前周日22时前提交，组内在当日18时前封版。当堂新意见安排在课堂后处理，不倒写为此前已完成。教师改变安排时，由组长同步更新计划、Project和提交清单。[R3]

'''+table(['课堂日期','邮件截止周日22时','提交材料','对应工作'],dates)+'''
组长使用学校邮箱提交，邮件主题为“SE2026-G04-作业名称”，按课程要求抄送助教和本组成员。收件、抄送、附件、来源及版本在发送前核对并保留发送记录。'''
team='''王晨担任组长兼项目负责人，负责范围、进度、风险、变更、阶段评审、集成和提交管理。刘景浩独立负责网页、社区及运维模块。双方使用AI完成实现，同时参加需求、设计、验证、文档和汇报，不以工具使用者之外的AI充当任务责任人。

'''+table(['成员','主要负责工作','共同与持续责任'],[
('王晨 32401211\n组长兼项目负责人','项目管理；后端骨架、认证、文章API、移动创作、交流激励、集成与性能部署验证','计划、可行性、SRS、设计评审、测试验收、文档与总结；协调跨模块交接和提交'),('刘景浩 32401205','双端阅读界面、网页Markdown、社区处理接口及网页、独立运维执行器与网页；功能安全兼容和回归统筹','原型、测试设计、报告、运维指南、PPT素材与归档；提交模块状态与风险；交叉复核')])+'''
共同任务明确双方工作

计划与可行性中，王晨整理范围、资源与关键技术，刘景浩整理排期、费用、操作与内容来源；双方确认结论。SRS中，王晨整理权限、业务和数据规则，刘景浩整理原型、操作与异常场景；双方确认可验收性。设计与联调中，各自完成模块，共同核对接口、数据和退出条件。测试与总结中，各自提供真实证据，刘景浩汇总材料，王晨核对技术限制和整体交付。

6.6由刘景浩统筹缺陷台账与回归；王晨修复其后端、移动和激励模块，刘景浩修复网页、社区及运维模块。交叉复核有独立任务，不代表把另一成员工作全部转交给复核者。

成员贡献按真实工作包、提交、问题关闭、文档与管理证据评价，不预设比例，也不按任务数量直接换算贡献。阶段评价检查及时性、成果质量、问题处理和协作；当前事实由双方确认后填写附录D。'''
quality='''质量保证从需求开始，REQ需求编号、TC用例编号、设计模块、WBS工作包、代码版本及证据相互关联。完成标准用检查场景、预期结果和应保存证据描述，避免只写“功能完成”。[R1，13.5]

'''+table(['检查点','检查内容','责任与产物'],[
('计划及评审','范围、期限、预算、依赖、资源与16项材料','双方共同确认，王晨组织；问题与处理记录'),('需求与设计','角色、状态、异常、数据、非功能和可验收性','各自编写与交叉复核；追踪表与基线'),('模块实现与接入','契约、权限、事务、草稿冲突和错误处理','任务负责人自查，另一人复核；单元接口与评审证据'),('测试、维护与交付','实测结果、缺陷关闭、恢复与版本一致','双方提供证据与修复；报告、维护单和发布清单')])
firstqa=section_first(8)
m=firstqa.find('| 场景 |');n=firstqa.find('\n单元检查',m)
quality+='\n'+firstqa[m:n].strip()+'\n\n'
quality+=table(['测试类别','对象与工具','责任与证据'],[
('单元','权限、状态、账本和筛选；JUnit及Spring测试','各模块负责人；代码版本与单元结果'),('接口、功能和集成','五类角色、双端与完整链路；Postman','刘景浩统筹、双方修复；集合、断言和缺陷'),('安全与兼容','越权、隐私、危险输入、图片、双端编辑','刘景浩；直接API、浏览器及开发者工具记录'),('性能与负载','查询、详情、发布与评论；JMeter','王晨；JTL、P95、错误率、吞吐与资源趋势'),('压力与稳定性','容量边界、持续运行与降载恢复','王晨；阶梯脚本、故障和一致性记录'),('部署与备份恢复','冷启动、部署重启、损坏备份、回退','王晨验证、刘景浩配合运维实现；恢复前后数据与执行日志')])
quality+='''
性能与恢复目标

以下是计划验收目标，最终结论必须注明执行环境、版本、数据与实际结果。同机施压需记录资源竞争；JMeter结果不能替代页面渲染和小程序体验检查。

'''+table(['场景','计划目标','必须保留的证据'],[
('30并发负载','合成200用户、2000文章、5000评论；60秒升载、运行10分钟；读P95≤1秒、写≤2秒，非预期错误率<1%','环境版本、脚本、原始JTL与统计结果'),('阶梯压力','10、30、50、100并发，各运行2分钟，升载30秒；100用于探测，不承诺此容量','吞吐、错误、CPU内存、失效点与数据一致性'),('稳定性与恢复','10并发运行30分钟；压力后降至10并发，2分钟内恢复；账本与数据不发生不一致','持续运行结果、降载恢复及数据库断言'),('备份恢复','数据库、图片及版本清单一致；拒绝损坏或不兼容备份；恢复前保护当前状态','备份哈希、恢复日志、关键记录及失败保护结果')])+'''
核心流程全部执行，阻塞与严重缺陷关闭。未关闭的一般问题列影响、责任人与期限。必须需求有实际验收结果；本机网页部署、业务重启、备份恢复有真实记录；代码、安装包、报告和测试版本一致。公网部署及小程序正式发布不作为本次课堂条件。'''
deploy=section_first(9)
deploy=deploy.replace('王晨负责运行架构、执行服务和技术恢复，刘景浩负责管理网页、操作手册和页面验收。','王晨负责集成环境、运行架构核对及部署恢复验收；刘景浩负责独立执行服务、运维网页、操作手册和模块修复；双方完成维护交接。')
deploy=deploy.replace('[R6','[R1').replace('10月11日前','10月12日前')
risk=section_first(10)
risk=risk.replace('[R6','[R1').replace('10月11日前','10月12日前')
risk+='''
资源与预算检查

Project将95项子任务、14项周管理和35项交叉复核同时建模。组长管理为持续责任，复核与文档也占用容量。当前排期模型中两位成员均未标为过载；这不是实际空闲时间的证明，周复盘仍需核对真实负担和考试等冲突。容量参数不作为贡献比例或人时，在公开甘特图只显示任务负责人。

'''+table(['预算项目','金额','依据与控制'],[('ChatGPT Plus订阅','20美元/月≈134.70元/月；4个月538.81元','1份订阅，2026年10月至2027年1月；按80美元×6.7351折算'),('豆包订阅','38元/月；4个月152.00元','1份订阅，单价采用小组提供价格'),('其他新增费用','0元','WorkBuddy及其他工具使用当前可用条件，设备与可选服务器复用；不列管理储备'),('项目预算合计','690.81元','两项订阅共用费用在2.2.1登记一次，其他任务不重复分摊；实际按账单登记')])+'''
ChatGPT Plus单价依据OpenAI官方定价页；汇率取中国银行2026年10月1日17:59:27美元中行折算价，100美元=673.51元，即1美元=6.7351元。总额先按未舍入美元金额计算，再保留两位小数，因此不使用单月近似人民币金额重复相乘。实际人民币扣款随结算汇率及账单变化；当前690.81元为计划估算。[R9][R10]

任务预计无法赶上交接即登记，由组长分析是否影响关键路径、测试和提交；另一成员承担可独立子任务或资料复核，重排后同步Project、周计划、风险和变更。'''
config=section_first(11).replace('[R6','[R1')
config=re.sub(r'阶段基线在评审通过.*?建立。','内部基线由双方评审并确认：需求对应10月23日，设计对应11月5日，实现对应11月27日，首轮内部验收对应12月4日。课堂评审基线在教师意见完成处理后另行标识，内部确认不得写成教师已批准。',config)
body=['# 城院帮项目计划\n\nG04 软件工程课程项目\n\n组长兼项目负责人 王晨 32401211；成员 刘景浩 32401205。\n\n计划基准日期 2026年10月1日；课程收尾 2027年1月6日。\n\n本计划用于计划阶段评审和后续执行，交付免费的实习与求职经验社区及管理维护能力。正文说明目标、安排和管理，附录提供完整周任务与验收依据。', '## 1 项目概况与目标\n\n'+intro,'## 2 范围与目标用户\n\n'+scope,'## 3 工程过程与阶段成果\n\n'+life,'## 4 系统架构与开发工具\n\n'+tools,'## 5 工作分解与进度计划\n\n'+wbs,'## 6 里程碑与提交安排\n\n'+milestones,'## 7 团队分工与成员评价\n\n'+team,'## 8 质量保证与异常处理\n\n'+quality,'## 9 部署运营与维护\n\n'+deploy,'## 10 风险资源与沟通管理\n\n'+risk,'## 11 配置变更与收尾管理\n\n'+config,'## 12 参考资料\n\n'+refs]
appendix='## 附录A 工作包时间责任与完成标准\n\n共用订阅费用在工作包2.2及子任务2.2.1登记690.81元，其余工作包预算0元；95项子任务在附录B按周展开。工作包负责人、复核人和依赖在Project使用同一数据。\n\n'
for s in range(1,9):
 rows=[]
 for t in tasks:
  if not t['code'].startswith(str(s)+'.'):continue
  rows.append([t['code']+' '+t['name'],datewindow(t)+'；'+cash(t),owner(t)+'；复核'+t['reviewer'],('前置'+','.join(t['preds'])+'；' if t['preds'] else '无前置；')+t['evidence']+'。完成标准：'+t['accept']])
 appendix+=f'### {s} '+['项目计划与评审','可行性研究','需求分析','总体与详细设计','编码与单元测试','综合测试','部署与维护','总结与收尾'][s-1]+'\n\n'+table(['工作包','日期与工期预算','任务负责人及复核','前置产物与完成标准'],rows)+'\n'
weektext='## 附录B 详细任务分解与周安排\n\n每行按所列子任务范围执行与检查，工作包整体完成标准见附录A。日为排期工作日，不列固定工作时段。子任务2.2.1登记全项目订阅预算690.81元；其余子任务、周管理与复核预算0元。双方每周按Project的R类任务检查对方产物，记录问题、修改和关闭结论；共同任务由双方确认。\n\n'
handoffs=['10月9日双方确认计划、图表和汇报；10月11日22时提交评审材料。关键技术验证接续至10月12日。','10月14日课堂意见在15日修订；可行性结论由双方确认，10月18日22时提交。','需求、权限、原型与验收条件形成同一基线；10月25日22时交SRS和需求PPT。','先按内部需求基线设计；10月28日已有意见进入变更台账，11月1日22时交已有修订。','11月4日课后处理需求复审意见；11月5日内部设计评审；11月6日开始实现；11月8日交总体设计。','提供认证和内容接口契约；双端页面模拟数据明确标注；11月15日交详细设计。','移动编辑与发布接口接通；刘景浩交付社区处理API和管理网页。设计修订材料11月22日提交。','运维执行器和网页完成交接；11月26日处理设计意见，27日完成真实API联调；29日交代码清单、用例及翻转PPT2。','测试使用固定数据和候选版本；双方按模块修复；12月4日内部验收，6日22时交报告。','12月6日已提交基线保留；后续报告复核单独修订，整理运行说明并检查冷启动；13日交已有测试修订。','双方完成维护与备份交接；17日处理测试课堂新意见；开始真实总结和偏差分析。','总结与总评素材相互核对；24日共同验收演练；27日22时交总结和总评PPT。','参与12月30日总评；31日准备1月3日提交基线，未完成尾项另列；元旦不排正常任务。','1月4日核对归档尾项；6日双方确认课程收尾和最终交付。']
WEEK_ACCEPT={
'5.1':('骨架与配置可构建，凭据独立','共享请求、错误格式与接口示例明确'),
'5.2':('登录验证码错误、过期、发信失败可处理','材料受控；审核理由、结果和角色权限正确'),
'5.3':('草稿不公开；正文、图片和版本关联正确','岗位专业公司组合筛选及分页正确'),
'5.4':('双端检索阅读按共享契约完成','空结果与失败可恢复；模拟数据标注清楚'),
'5.5':('Markdown预览、图片和网盘链接可用','保存失败保留正文，危险脚本受控，草稿可恢复'),
'5.6':('手机图文与图片上传可独立完成','失败图片可重试，草稿恢复不覆盖复杂正文'),
'5.7':('生命周期接口及冲突检测可用','双端管理一致；管理下架不能自行恢复'),
'5.8':('评论图片、主页联系、收藏有认证边界','余额不负、不重复奖励、不自投；荣誉及通知一致'),
'5.9':('处理API复用认证，状态理由审计齐备','网页定位原文评论；处理、申诉和通知联通'),
'5.10':('独立执行器鉴权、目标白名单及任务互斥可用','四类操作留痕；数据库图片一致，恢复前保护当前状态'),
'5.11':('状态、日志及备份列表权限正确','网页确认目标和影响，失败及停机仍能查结果'),
'5.12':('双端业务及运维接真实API，关闭模拟替代缺口','构建、源码、代码清单和用例同版本归档'),
'6.1':('五类角色和完整核心链路实际执行','失败可复现，输入数据断言关联缺陷'),
'6.2':('固定数据环境执行30并发负载场景','核对P95、错误率、吞吐和资源原始结果'),
'6.3':('阶梯10至100并发探测边界，非容量承诺','10并发30分钟；降载恢复及账本一致性有记录'),
'6.4':('直接API越权失败，隐私与危险输入受控','浏览器和开发者工具核心体验一致，设备限制列明'),
'6.5':('指南冷启动及网页四类运维操作实际执行','损坏备份被拒；回退与恢复前后数据一致'),
'6.6':('双方按模块修复，记录版本和缺陷指派','交叉回归原及相邻场景，严重缺陷关闭'),
'6.7':('实际结果与目标分开，环境版本证据可追溯',''),
'6.8':('已提交测试基线保留，复核结论另记','报告修订与课堂汇报证据一致'),
'8.4':('12月31日形成1月3日待提交基线','1月4日核对尾项，以补充记录归档'),
}
for w,handoff in zip(weeks,handoffs):
 a=date.fromisoformat(w['start']);b=date.fromisoformat(w['finish']);weektext+=f'### 第{w["number"]}周 {a:%m/%d}—{b:%m/%d}\n\n本周目标：{w["goal"]}。\n\n交接与提交：{handoff}\n\n'
 rows=[]
 for c in w['subtasks']:
  q=idx[c];part=int(c.rsplit('.',1)[1])-1
  standard=WEEK_ACCEPT.get(q['parent'])
  criterion=standard[part] if standard else ('形成'+q['name']+'对应产物；按附录A核对内容与未决问题' if part==0 else '完成'+q['name']+'；记录交叉复核或双方确认结论')
  name=q['name'];window=datewindow(q);duration=cash(q)
  if datetime.fromisoformat(q['start']).date()<a or datetime.fromisoformat(q['finish']).date()>b:
   sa=max(datetime.fromisoformat(q['start']).date(),a);fb=min(datetime.fromisoformat(q['finish']).date(),b)
   while fb.weekday()>=5 or fb==date(2027,1,1):fb-=timedelta(days=1)
   while sa.weekday()>=5 or sa==date(2027,1,1):sa+=timedelta(days=1)
   window=sa.strftime('%m/%d')+('—'+fb.strftime('%m/%d') if sa!=fb else '')
   n=sum((sa+timedelta(days=j)).weekday()<5 and sa+timedelta(days=j)!=date(2027,1,1) for j in range((fb-sa).days+1));duration=f'本周{n}日；全任务{q["duration"]:g}日；新增0元'
   beginning=datetime.fromisoformat(q['start']).date()>=a
   name=('本周启动：' if beginning else '本周收尾：')+name
   if c=='2.1.2':criterion='确认邮件渠道、运维独立执行和备份范围的配置条件' if beginning else '执行样例核查并记录真实结果、失败与替代措施'
   if c=='4.3.2':criterion='形成图片上传、草稿与跨端版本冲突流程草案' if beginning else '核对复杂正文保留和失败恢复，完成详细设计复核'
  rows.append([c,name+'；'+window+'；'+duration,owner(q),','.join(q['preds']) if q['preds'] else '无；按约定输入启动',criterion+'。证据：'+q['evidence']])
 m=w['management']
 weektext+='持续管理 '+m['code']+'：'+datewindow(m)+'；'+cash(m)+'。负责人王晨，刘景浩提供状态。'+m['accept']+'；保存周会议及变更记录。\n\n'
 weektext+=table(['编号','当周具体任务','任务负责人','前置与交接','产物与完成标准'],rows)+'\n'
rubric=[('演示PPT及效果','豆包PPT提示词和配套图表','按正式素材生成后检查排版、来源和附录'),('按进度提交阶段成果','正文第3、6节；附录B','提交记录以实际发送为准'),('可行性与关键技术','可行性报告；2.1—2.3','方案与真实验证结果分开'),('计划模板适合项目','正文12节及四项附录','目标、分工、排期和管理成体系'),('成员贯穿全程与甘特图','第7节；附录B；MPP','两人有开发、文档、验证和评审责任'),('使用Project且图准确','MPP；原生甘特图PDF；核查CSV','48工作包、95子任务日期工期逐项核对'),('明确合理分工','第7节；附录A、B','管理、开发、文档、共同任务明确'),('WBS与瀑布阶段','第3、5节；Project层级','8工程阶段、工作包和子任务，反馈单列'),('逐项完成时间估计','附录A、B；任务CSV','起止日期、明确工期和依赖，无问号'),('逐项指定任务负责人','附录A、B；Project任务表','学生为负责人，共同任务列两人'),('适配课程进度','第6节；周计划和K类节点','周日提交与课后意见处理分开'),('整体与逐项预算','第10节；附录A、B','两项订阅4个月合计690.81元；2.2.1登记一次，其他0元'),('文档配置管理','第11节；Git文档版本资料','正文图表源稿与导出记录入同一基线'),('负荷不均或过载说明','第10节；原生资源核查','开发文档管理复核同时建模，真实负担周检'),('会议及有效记录','附录D；会议模板','双方填写真实日期、决议、行动及完成证据'),('当前成员绩效评价','第7节；附录D','按真实成果确认，不代填会议或评价分数')]
review='## 附录C 计划阶段16项评审对照\n\n'+table(['条目','教师检查要求','材料位置','核查口径'],[[str(i),*r] for i,r in enumerate(rubric,1)])
meeting='''## 附录D 会议与当前成员评价记录

阶段会议记录应包括实际时间、参与人、工作包完成证据、问题、风险变更、决议以及带负责人和期限的行动项。每周由组长组织，刘景浩提供模块状态，双方确认结论。模板不等同于已召开会议。

'''+table(['记录字段','填写内容'],[('实际日期与参与人','按真实会议填写'),('讨论与依据','任务编号、产物版本、问题或教师意见'),('决议','范围、依赖、分工或排期改变及理由'),('行动项','负责人、完成期限、验收证据与关闭状态')])+'''
当前阶段成员评价由两人根据实际计划、文档、图表、管理和协作成果确认。对王晨检查项目组织、范围进度与技术材料；对刘景浩检查排期资料、页面及文档工作和协作。描述采用“完成了哪些成果、质量与及时性、存在的问题、后续改进”，不预填贡献比例、会议出席或成绩。
'''
body +=[appendix,weektext,review,meeting]
plan='\n\n'.join(body)
plan=plan.replace('新版','当前').replace('旧版','历史').replace('[R6，第5—20页]','[R1，13.1至13.3]')
(OUT/'城院帮项目计划G04.md').write_text(plan,encoding='utf8')
(OUT/'详细周计划G04.md').write_text(weektext,encoding='utf8')
(OUT/'评审表16项对照.md').write_text(review,encoding='utf8')
(OUT/'会议与成员评价记录模板.md').write_text(meeting,encoding='utf8')

# Preserve the detailed feasibility narrative and update only its affected interfaces and references.
fea=(ROOT/'archive/项目计划历史/20261001/deliverables/G04-project-plan-20261001/城院帮可行性分析报告G04.md').read_text(encoding='utf8')
fea=fea.replace('项目计划第2、3、6、11至17节及工作包2.1至2.3','项目计划第2、4、5、7至11节及工作包2.1至2.3')
fea=fea.replace('48个工作包','48个工作包、95项子任务')
fea=re.sub(r'杨桦[．.\s]*','',fea)
fea=fea.replace('第13章 软件项目管理','课程课件 第13章 软件项目管理').replace('第1章 软件工程学概述','课程课件 第1章 软件工程学概述')
fea=fea.replace('新版','当前').replace('旧版','历史')
fea=fea.replace('本文附录C逐项对应16项检查内容','项目计划附录C逐项对应16项检查内容')
fea=fea.replace('https://jmeter.apache.org/usermanual/generating-dashboard ','https://jmeter.apache.org/usermanual/generating-dashboard.html ').replace('https://jmeter.apache.org/usermanual/glossary ','https://jmeter.apache.org/usermanual/glossary.html ')
fea=fea.replace('已拥有的电脑和可用订阅作为存量资源，不再虚构采购。','已拥有的电脑作为存量资源；ChatGPT Plus和豆包订阅按2026年10月至2027年1月共4个月计入项目预算。WorkBuddy用于小程序开发，采用当前可用环境，不另列新增采购。')
fea=fea.replace('| 直接新增预算 | 48个工作包、95项子任务均0元 | 不计划新增采购，不把人时转换为金额 |\n| 管理储备 | 200元，整体新增现金上限200元 | 临时必要支出先登记，超支前修订计划 |','| ChatGPT Plus订阅 | 20美元/月约134.70元；4个月538.81元 | 1份订阅，80美元×6.7351；官方定价与中行折算价 |\n| 豆包订阅 | 38元/月；4个月152.00元 | 1份订阅，小组提供单价 |\n| 其他新增费用 | 0元 | 设备、服务器和当前可用工具复用，不列管理储备 |\n| 项目预算合计 | 690.81元 | 共用费用在2.2.1登记一次；其他任务不重复分摊 |')
fea=fea.replace('条件结论为经济上可承担。若出现必须购买且超过储备的工具，优先选择能够产出同类成果的替代方案；确需新增采购时再评估对预算和交付的影响。','条件结论为经济上可承担。预算只包含两项订阅，共690.81元。ChatGPT Plus人民币金额按中国银行2026年10月1日17:59:27中行折算价1美元=6.7351元估算，先计算80美元的总额再保留两位小数；实际按结算汇率和账单登记。其他工具采用已有条件，如出现新增费用须通过预算变更再评估。[9][10]')
fea=fea.replace('控制新增支出200元上限，核实工具可用性','订阅预算690.81元，核实WorkBuddy及其他工具可用性')
fea+='\n\n[9] OpenAI官方文档 Pricing，https://learn.chatgpt.com/docs/pricing ，2026年10月1日核对，ChatGPT Plus每月20美元。\n\n[10] 中国银行 外汇牌价，https://www.boc.cn/sourcedb/whpj/ ，2026年10月1日17:59:27，美元中行折算价673.51元/100美元。\n'
(OUT/'城院帮可行性分析报告G04.md').write_text(fea,encoding='utf8')

def setfont(run,size=12,bold=False):
 run.font.name='Microsoft YaHei';run.font.size=Pt(size);run.font.bold=bold
 run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
def docbase(title):
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=s.right_margin=Cm(1.8);s.top_margin=Cm(1.9);s.bottom_margin=Cm(1.7)
 for name,size in [('Normal',12),('Title',22),('Subtitle',14),('Heading 1',16),('Heading 2',14),('Heading 3',12)]:
  st=d.styles[name];st.font.name='Microsoft YaHei';st.font.size=Pt(size);st.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0);st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑');st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.13
 d.styles['Normal'].paragraph_format.space_after=Pt(6);d.core_properties.author='G04';d.core_properties.title=title
 p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 setfont(p.add_run('第 '),10.5);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld);setfont(p.add_run(' 页'),10.5)
 return d
def grid(d,heads,rows):
 t=d.add_table(rows=1,cols=len(heads));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths={2:[4.0,14.0],3:[3.4,7.1,7.5],4:[4.2,3.0,3.4,7.4],5:[1.5,4.3,2.7,2.2,7.3]}[len(heads)]
 widths=[x*17.4/18 for x in widths]
 for i,w in enumerate(widths):t.columns[i].width=Cm(w)
 for i,c in enumerate(t.rows[0].cells):c.text=heads[i];c.width=Cm(widths[i]);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2');c._tc.get_or_add_tcPr().append(shade)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  for i,(c,text) in enumerate(zip(t.add_row().cells,row)):c.text=text;c.width=Cm(widths[i])
 for ri,row in enumerate(t.rows):
  no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
  for c in row.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.08;p.paragraph_format.keep_with_next=(ri==0)
    for run in p.runs:setfont(run,10.5,ri==0)
   pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'BBC6CF');borders.append(b)
   pr.append(borders)
 return t
def render_md(text,title,path,is_plan=False):
 d=docbase(title);p=d.add_paragraph(title,'Title');d.add_paragraph('软件工程课程项目 G04','Subtitle');d.add_paragraph('组长兼项目负责人 王晨 32401211\n成员 刘景浩 32401205');d.add_paragraph('2026年10月1日')
 d.add_paragraph('实习与求职经验分享社区\n免费阅读 非现金激励 自愿站外联系');d.add_paragraph('本文用于项目计划阶段评审及后续执行。正文说明范围、责任、进度和管理，附录与图表提供逐任务检查依据。' if is_plan else '本报告从关键技术、资源费用、操作、进度和风险判断项目可行条件，并与项目计划保持一致。')
 d.add_page_break();p=d.add_paragraph();setfont(p.add_run('目录'),16,True);p=d.add_paragraph();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u');p._p.append(f);d.add_page_break()
 if is_plan:text=text[text.find('## 1 '):]
 if not is_plan:
  # Existing source is a plain export with numbered section lines.
  text=re.sub(r'^([1-9]\d*) ([^\n]+)$',r'## \1 \2',text,flags=re.M)
  pos=text.find('## 1 ');text=text[pos:]
 lines=text.splitlines();i=0
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('# ') or line in ['目录']:i+=1;continue
  if line.startswith('## '):
   h=line[3:]
   p=d.add_paragraph(h,'Heading 1')
   if is_plan and (h.startswith('附录') or h.startswith('12 参考')):p.paragraph_format.page_break_before=True
   i+=1;continue
  if line.startswith('### '):
   h=line[4:]
   p=d.add_paragraph(h,'Heading 2')
   i+=1;continue
  if line.startswith('|'):
   block=[]
   while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i].strip());i+=1
   rows=[[s.strip() for s in l.strip('|').split('|')] for l in block];rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',v) for v in r)]
   if rows:grid(d,rows[0],rows[1:])
   gap=d.add_paragraph();gap.paragraph_format.space_before=Pt(0);gap.paragraph_format.space_after=Pt(2);gap.paragraph_format.line_spacing=Pt(1);gap.paragraph_format.keep_with_next=True;setfont(gap.add_run(),1);continue
  if line.startswith('!['):
   m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
   if m:
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(OUT/m.group(2)),width=Cm(17.4));cap=d.add_paragraph(m.group(1));cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
   i+=1;continue
  p=d.add_paragraph(re.sub(r'\*\*(.*?)\*\*',r'\1',line))
  if is_plan and line.startswith(('本周目标','交接与提交','持续管理')):p.paragraph_format.keep_with_next=True
  i+=1
 d.save(path)
render_md(plan,'城院帮项目计划',OUT/'城院帮项目计划G04.docx',True)
render_md(fea,'城院帮可行性分析报告',OUT/'城院帮可行性分析报告G04.docx')
print(json.dumps({'plan_sections':12,'packages':48,'subtasks':95,'weeks':14,'tables':len(Document(OUT/'城院帮项目计划G04.docx').tables)},ensure_ascii=False))


