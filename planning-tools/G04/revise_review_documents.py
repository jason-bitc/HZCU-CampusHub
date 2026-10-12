from pathlib import Path
import shutil,re,copy,json
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'deliverables/G04-项目计划'
HISTORY=ROOT/'archive/评审修订前资料/20261006';HISTORY.mkdir(parents=True,exist_ok=True)
names=['城院帮项目计划G04.docx','城院帮项目计划G04.pdf','城院帮项目计划G04.md','详细周计划G04.md','豆包PPT制作提示词.md','评审表16项对照.md','会议与成员评价记录模板.md','请先阅读.md','交付检查记录.md']
for name in names:
    if not (HISTORY/name).exists():shutil.copy2(OUT/name,HISTORY/name)
src=ROOT/'deck6/exports/城院帮项目计划（G04）.pptx'
if not (HISTORY/src.name).exists():shutil.copy2(src,HISTORY/src.name)

template='报告结构依据课程的软件工程阶段、第13章项目管理方法及计划阶段评审要求组织。正文说明目标、范围、责任、排期、质量、部署维护、风险预算与配置管理；工作包、周计划、评审对照及阶段记录附录提供逐项检查依据。'
meeting='''## 附录D 会议与当前成员评价记录

截至2026年10月6日，双方实际会议与当前阶段成员评价尚未确认。以下为待填写记录表，不作为已完成会议或评价的证据。完成真实记录后，由双方核对日期、成果、决议及行动项，并同步到Word、PDF及PPT；无需预设贡献比例或编造评价分数。

### 会议记录确认表

| 事项 | 实际记录 |
| --- | --- |
| 实际日期和参与人 | 尚未确认 |
| 讨论依据与产物版本 | 待填写任务编号、成果文件或教师意见 |
| 共同决定 | 待填写双方实际决定，保留未解决事项 |
| 行动与任务负责人 | 待填写负责人、期限、产物和完成条件 |
| 处理与复核结果 | 待填写证据位置、实际完成日期及双方确认 |

### 当前阶段成员评价确认表

| 成员 | 实际成果与证据 | 质量及时性及改进 | 双方确认 |
| --- | --- | --- | --- |
| 王晨 32401211 | 尚未确认；待填写实际产物及个人承担部分 | 待填写已确认成果的质量、及时性和后续改进 | 待确认 |
| 刘景浩 32401205 | 尚未确认；待填写实际产物及个人承担部分 | 待填写已确认成果的质量、及时性和后续改进 | 待确认 |

评价按真实成果、问题关闭及协作证据形成；计划中的职责不等同于已经完成的工作。阶段15、16项仍须由双方补齐真实记录。
'''

def text_replace(doc,old,new):
    for p in list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]:
        if old in p.text:
            p.text=p.text.replace(old,new)
            for run in p.runs:
                run.font.name='Microsoft YaHei';run.font.size=Pt(10.5 if p._p.getparent().tag==qn('w:tc') else 12)
                run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')

def para_before(doc,element,text,style='Normal'):
    p=doc.add_paragraph(text,style);element.addprevious(p._p);return p

def table_before(doc,element,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,s in zip(t.rows[0].cells,headers):c.text=s
    for row in rows:
        for c,s in zip(t.add_row().cells,row):c.text=s
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri,row in enumerate(t.rows):
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
            for side in ['top','bottom','left','right']:
                q=OxmlElement('w:'+side);q.set(qn('w:val'),'single');q.set(qn('w:sz'),'4');q.set(qn('w:color'),'D9D9D9');bd.append(q)
            tcpr.append(bd)
            if ri==0:
                q=OxmlElement('w:shd');q.set(qn('w:fill'),'E8EDF2');tcpr.append(q)
            for p in c.paragraphs:
                p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.08
                p.paragraph_format.keep_with_next=(ri==0)
                for r in p.runs:
                    r.font.name='Microsoft YaHei';r.font.size=Pt(10.5);r.font.bold=(ri==0)
                    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    element.addprevious(t._tbl);return t

doc=Document(HISTORY/'城院帮项目计划G04.docx')
text_replace(doc,'13.2成本估算','13.2工作量估算')
text_replace(doc,'13.2 成本估算','13.2 工作量估算')
text_replace(doc,'10/11前验证WSL2','10/12前完成WSL2')
text_replace(doc,'10/11 前验证 WSL2','10/12 前完成 WSL2')
text_replace(doc,'10/11前验证 WSL2','10/12前完成 WSL2')
text_replace(doc,'豆包PPT提示词和配套图表','实际48页PPT与配套图表')
text_replace(doc,'双方填写真实日期、决议、行动及完成证据','当前尚未确认；双方补齐真实日期、决议、行动及完成证据')
text_replace(doc,'按真实成果确认，不代填会议或评价分数','当前尚未确认；按双方真实成果确认评价，不预设比例')
heading=next(p for p in doc.paragraphs if p.text=='2 范围与目标用户')
if not any(p.text==template for p in doc.paragraphs):para_before(doc,heading._p,template)
dhead=next(p for p in doc.paragraphs if p.style.name=='Heading 1' and (p.text.startswith('附录D ') or p.text.startswith('附录 D ')))
body=doc._element.body
for e in list(body)[list(body).index(dhead._p)+1:]:
    if e.tag!=qn('w:sectPr'):body.remove(e)
end=body.find(qn('w:sectPr'))
para_before(doc,end,'截至2026年10月6日，双方实际会议与当前阶段成员评价尚未确认。以下为待填写记录表，不作为已完成会议或评价的证据。完成真实记录后双方核对，并同步到报告与PPT。')
para_before(doc,end,'会议记录确认表','Heading 2')
table_before(doc,end,['事项','实际记录'],[
    ['实际日期和参与人','尚未确认'],['讨论依据与产物版本','待填写任务编号、成果文件或教师意见'],['共同决定','待填写双方实际决定，保留未解决事项'],['行动与任务负责人','待填写负责人、期限、产物和完成条件'],['处理与复核结果','待填写证据位置、实际完成日期及双方确认']],[4.0,13.4])
para_before(doc,end,'当前阶段成员评价确认表','Heading 2')
table_before(doc,end,['成员','实际成果与证据','质量及时性及改进','双方确认'],[
    ['王晨\n32401211','尚未确认；待填写实际产物及个人承担部分','待填写成果质量、及时性及后续改进','待确认'],['刘景浩\n32401205','尚未确认；待填写实际产物及个人承担部分','待填写成果质量、及时性及后续改进','待确认']],[2.6,6.0,6.2,2.6])
para_before(doc,end,'评价按真实成果、问题关闭及协作证据形成，计划中的职责不等同于已完成工作。第15、16项须双方补齐真实记录；不预设贡献比例或编造评价分数。')
all_text='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells])
subs=json.loads((OUT/'图表与排期/子任务数据.json').read_text(encoding='utf8'))
assert len(doc.tables)>30 and all(q['code'] in all_text for q in subs)
doc.save(OUT/'城院帮项目计划G04.docx')

p=OUT/'城院帮项目计划G04.md';s=p.read_text(encoding='utf8')
s=s.replace('13.2成本估算','13.2工作量估算').replace('10/11前验证WSL2','10/12前完成WSL2')
if template not in s:s=s.replace('## 2 范围与目标用户',template+'\n\n## 2 范围与目标用户')
s=s.replace('豆包PPT提示词和配套图表','实际48页PPT与配套图表')
s=s.replace('双方填写真实日期、决议、行动及完成证据','当前尚未确认；双方补齐真实日期、决议、行动及完成证据')
s=s.replace('按真实成果确认，不代填会议或评价分数','当前尚未确认；按双方真实成果确认评价，不预设比例')
s=re.sub(r'## 附录D .*',meeting,s,flags=re.S)
p.write_text(s,encoding='utf8')
(OUT/'会议与成员评价记录模板.md').write_text(meeting,encoding='utf8')
p=OUT/'评审表16项对照.md';s=p.read_text(encoding='utf8')
s=s.replace('豆包PPT提示词和配套图表','实际48页PPT与配套图表')
s=s.replace('双方填写真实日期、决议、行动及完成证据','当前尚未确认；双方补齐真实日期、决议、行动及完成证据')
s=s.replace('按真实成果确认，不代填会议或评价分数','当前尚未确认；按双方真实成果确认评价，不预设比例')
p.write_text(s,encoding='utf8')

p=OUT/'豆包PPT制作提示词.md';s=p.read_text(encoding='utf8')
s=s.replace('岗位查经验→同专业或公司筛选→阅读收藏→认证评论→可选站外联系','岗位查经验→同专业或公司筛选→免费阅读→认证后收藏评论→可选站外联系')
s=s.replace('工作包整体完成标准见附录A','工作包整体完成标准见项目计划报告附录A')
heading='## 评审检查后的制作约束'
if heading in s:s=s[:s.index(heading)].rstrip()
s+='\n\n'+heading+'\n\n保持48页结构和完整95项子任务；正式PPT文件名为城院帮项目计划G04.pptx。正文说明文字提高可读性，保留完整电子附录，不裁切工程图的标题、节点或图例。5.9总工期3日；5.10本周2日、总工期3日，11月23日收尾。第12页图注为设计与复审网络，不使用错误事件范围；第14页图注为联调与验证网络，不使用错误事件范围。\n\n第19页补11月4日SRS修订与复审、11月25日设计修订与翻转PPT1，并明确2.1至2.3是早期条件核查与可行性评审；完整系统验收在后续实现测试。资源结论为当前模型两人无过载，包含开发、文档、管理和复核；真实负担每周核对。\n\n用户已确认实际会议和当前成员评价尚未确认。第20页及评审第15、16项明确标为待双方确认，使用待填写记录表，不虚构日期、决议、出席、完成成果或比例。评审对照中所有章节与附录定位写明项目计划报告；PPT位置用页码，第1项指实际PPT，不再用提示词充当已生成PPT的证据。\n\n第3页需求依据引用本组校友圈观察与需求讨论；方法依据引用课件。第21页补教材结构参考与本组观察来源，完整链接放讲者备注。正式提交时必须包含实际PPT。\n'
p.write_text(s,encoding='utf8')
p=OUT/'阶段事实状态.json'
p.write_text(json.dumps({'确认日期':'2026-10-06','来源':'用户直接回复','实际会议记录':'尚未确认','当前成员绩效评价':'尚未确认','说明':'任务分工与生成文件不自动等同个人实际成果；待双方据事实填写。'},ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'plan_docx_updated':True,'meeting_status':'尚未确认','performance_status':'尚未确认','budget':690.81,'source_deck_preserved':str(HISTORY/src.name)},ensure_ascii=False))
